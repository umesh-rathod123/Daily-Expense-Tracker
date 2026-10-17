import csv
import json
import shutil
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import date, datetime, timedelta
from collections import defaultdict
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from openpyxl import load_workbook, Workbook

from database import initialize_database, EXCEL_FILE, DATA_DIR, HEADERS

# Daily Expense Tracker - expanded version
# Existing expenses.xlsx and its records are preserved. Extra settings are stored
# in separate JSON files inside the project's data folder.

initialize_database()

APP_DIR = Path(__file__).resolve().parent
BUDGET_FILE = DATA_DIR / "budgets.json"
RECURRING_FILE = DATA_DIR / "recurring_transactions.json"
DEFAULT_CATEGORIES = [
    "Food", "Travel", "Shopping", "Bills", "Education", "Health",
    "Entertainment", "Rent", "Fees", "Salary", "Business", "Other"
]

def ensure_support_files():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    for path, default in ((BUDGET_FILE, {}), (RECURRING_FILE, [])):
        if not path.exists():
            path.write_text(json.dumps(default, indent=2), encoding="utf-8")

def load_json(path, fallback):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return fallback

def save_json(path, data):
    ensure_support_files()
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")

ensure_support_files()

def format_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text

def parse_date(value):
    return datetime.strptime(format_date(value), "%Y-%m-%d").date()

def read_transactions():
    if not EXCEL_FILE.exists():
        initialize_database()
    workbook = load_workbook(EXCEL_FILE)
    try:
        if "Transactions" not in workbook.sheetnames:
            raise ValueError('The Excel workbook does not contain a "Transactions" sheet.')
        sheet = workbook["Transactions"]
        records = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            records.append({
                "ID": row[0],
                "Date": format_date(row[1]),
                "Type": str(row[2] or "Expense"),
                "Category": str(row[3] or ""),
                "Amount": float(row[4] or 0),
                "Description": str(row[5] or ""),
            })
        return records
    finally:
        workbook.close()

def write_transactions(records):
    """Rewrite only the Transactions sheet, preserving any other workbook sheets."""
    workbook = load_workbook(EXCEL_FILE)
    try:
        if "Transactions" not in workbook.sheetnames:
            sheet = workbook.create_sheet("Transactions")
            sheet.append(HEADERS)
        else:
            sheet = workbook["Transactions"]
        # Data rows only; keep the existing header and any other sheets.
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        for r in records:
            sheet.append([r["ID"], r["Date"], r["Type"], r["Category"], r["Amount"], r["Description"]])
        workbook.save(EXCEL_FILE)
    finally:
        workbook.close()

def next_id(records):
    nums = [int(r["ID"]) for r in records if str(r["ID"]).isdigit()]
    return max(nums, default=0) + 1

def validate_transaction(txn_type, category, amount_text, txn_date, description):
    category = category.strip()
    description = description.strip()
    if txn_type not in ("Income", "Expense"):
        raise ValueError("Choose Income or Expense.")
    if not category:
        raise ValueError("Category cannot be empty.")
    try:
        amount = float(amount_text)
    except (TypeError, ValueError):
        raise ValueError("Enter a valid numeric amount.")
    if amount <= 0:
        raise ValueError("Amount must be greater than zero.")
    if not (amount < float("inf")):
        raise ValueError("Enter a finite amount.")
    try:
        parsed = datetime.strptime(txn_date.strip(), "%Y-%m-%d").date()
    except ValueError:
        raise ValueError("Date must use YYYY-MM-DD format.")
    return category, round(amount, 2), parsed.isoformat(), description

def get_categories():
    names = set(DEFAULT_CATEGORIES)
    for r in read_transactions():
        if r["Category"].strip():
            names.add(r["Category"].strip())
    return sorted(names, key=str.lower)

def totals(records):
    income = sum(r["Amount"] for r in records if r["Type"] == "Income")
    expenses = sum(r["Amount"] for r in records if r["Type"] == "Expense")
    return income, expenses, income - expenses

def period_records(records, start, end):
    result = []
    for r in records:
        try:
            d = parse_date(r["Date"])
        except ValueError:
            continue
        if start <= d <= end:
            result.append(r)
    return result

def current_month_bounds():
    today = date.today()
    return date(today.year, today.month, 1), today

def money(value):
    return f"₹{value:,.2f}"

# ---------------- Main window ----------------
root = tk.Tk()
root.title("Daily Expense Tracker")
root.geometry("1180x780")
root.minsize(980, 650)

style = ttk.Style()
try:
    style.theme_use("clam")
except tk.TclError:
    pass

main = ttk.Frame(root, padding=10)
main.pack(fill="both", expand=True)

ttk.Label(main, text="Daily Expense Tracker", font=("Arial", 20, "bold")).pack(anchor="w")
ttk.Label(main, text=f"Excel data: {EXCEL_FILE}", foreground="#555").pack(anchor="w", pady=(0, 8))

summary = ttk.Frame(main)
summary.pack(fill="x", pady=4)
income_label = ttk.Label(summary, text="Total Income: ₹0.00", font=("Arial", 11, "bold"))
expense_label = ttk.Label(summary, text="Total Expenses: ₹0.00", font=("Arial", 11, "bold"))
balance_label = ttk.Label(summary, text="Current Balance: ₹0.00", font=("Arial", 11, "bold"))
savings_label = ttk.Label(summary, text="Savings Rate: N/A", font=("Arial", 11, "bold"))
for widget in (income_label, expense_label, balance_label, savings_label):
    widget.pack(side="left", padx=(0, 18), pady=4)

insights_var = tk.StringVar(value="Spending insights will appear here.")
ttk.Label(main, textvariable=insights_var, wraplength=1100).pack(anchor="w", pady=4)

# Notebook keeps all controls accessible without a very tall window.
tabs = ttk.Notebook(main)
tabs.pack(fill="both", expand=True, pady=6)
transactions_tab = ttk.Frame(tabs, padding=8)
reports_tab = ttk.Frame(tabs, padding=8)
budgets_tab = ttk.Frame(tabs, padding=8)
recurring_tab = ttk.Frame(tabs, padding=8)
tools_tab = ttk.Frame(tabs, padding=8)
tabs.add(transactions_tab, text="Transactions")
tabs.add(reports_tab, text="Reports & Charts")
tabs.add(budgets_tab, text="Budgets")
tabs.add(recurring_tab, text="Recurring")
tabs.add(tools_tab, text="Export & Backup")

# ---------------- Add/edit transaction form ----------------
form = ttk.LabelFrame(transactions_tab, text="Add Transaction", padding=8)
form.pack(fill="x", pady=(0, 8))
type_var = tk.StringVar(value="Expense")
category_var = tk.StringVar()
amount_var = tk.StringVar()
txn_date_var = tk.StringVar(value=date.today().isoformat())
description_var = tk.StringVar()
editing_id = tk.StringVar(value="")

def refresh_category_choices():
    category_combo["values"] = get_categories()

ttk.Label(form, text="Type").grid(row=0, column=0, sticky="w", padx=4, pady=3)
ttk.Combobox(form, textvariable=type_var, values=["Income", "Expense"], state="readonly", width=12).grid(row=0, column=1, padx=4, pady=3)
ttk.Label(form, text="Category").grid(row=0, column=2, sticky="w", padx=4, pady=3)
category_combo = ttk.Combobox(form, textvariable=category_var, values=DEFAULT_CATEGORIES, width=18)
category_combo.grid(row=0, column=3, padx=4, pady=3)
ttk.Label(form, text="Amount (₹)").grid(row=0, column=4, sticky="w", padx=4, pady=3)
ttk.Entry(form, textvariable=amount_var, width=14).grid(row=0, column=5, padx=4, pady=3)
ttk.Label(form, text="Date (YYYY-MM-DD)").grid(row=1, column=0, sticky="w", padx=4, pady=3)
ttk.Entry(form, textvariable=txn_date_var, width=16).grid(row=1, column=1, padx=4, pady=3)
ttk.Label(form, text="Description").grid(row=1, column=2, sticky="w", padx=4, pady=3)
ttk.Entry(form, textvariable=description_var, width=35).grid(row=1, column=3, columnspan=2, sticky="ew", padx=4, pady=3)

def clear_form():
    editing_id.set("")
    type_var.set("Expense")
    category_var.set("")
    amount_var.set("")
    txn_date_var.set(date.today().isoformat())
    description_var.set("")
    save_button.config(text="Add Transaction")

def save_transaction():
    try:
        category, amount, txn_date, description = validate_transaction(
            type_var.get(), category_var.get(), amount_var.get(),
            txn_date_var.get(), description_var.get()
        )
        records = read_transactions()
        target_id = editing_id.get()
        if target_id:
            found = False
            for r in records:
                if str(r["ID"]) == target_id:
                    r.update({"Type": type_var.get(), "Category": category, "Amount": amount,
                              "Date": txn_date, "Description": description})
                    found = True
                    break
            if not found:
                raise ValueError("Selected transaction was not found.")
            write_transactions(records)
            messagebox.showinfo("Updated", "Transaction updated successfully.")
        else:
            records.append({"ID": next_id(records), "Date": txn_date, "Type": type_var.get(),
                            "Category": category, "Amount": amount, "Description": description})
            # Append rather than rewriting existing data on ordinary additions.
            wb = load_workbook(EXCEL_FILE)
            try:
                ws = wb["Transactions"]
                ws.append([records[-1]["ID"], txn_date, type_var.get(), category, amount, description])
                wb.save(EXCEL_FILE)
            finally:
                wb.close()
            messagebox.showinfo("Saved", "Transaction saved successfully.")
        clear_form()
        refresh_category_choices()
        refresh_all()
    except (ValueError, OSError, PermissionError) as exc:
        messagebox.showerror("Cannot save transaction", str(exc))

save_button = ttk.Button(form, text="Add Transaction", command=save_transaction)
save_button.grid(row=1, column=5, padx=4, pady=3)
ttk.Button(form, text="Clear Form", command=clear_form).grid(row=0, column=6, rowspan=2, padx=4)

# ---------------- Search and transaction table ----------------
filters = ttk.LabelFrame(transactions_tab, text="Search & Filter", padding=7)
filters.pack(fill="x", pady=(0, 8))
search_var = tk.StringVar()
filter_type_var = tk.StringVar(value="All")
filter_category_var = tk.StringVar(value="All")
filter_start_var = tk.StringVar()
filter_end_var = tk.StringVar()
min_amount_var = tk.StringVar()
max_amount_var = tk.StringVar()

ttk.Label(filters, text="Search").grid(row=0, column=0, padx=3, pady=3)
ttk.Entry(filters, textvariable=search_var, width=17).grid(row=0, column=1, padx=3)
ttk.Label(filters, text="Type").grid(row=0, column=2, padx=3)
ttk.Combobox(filters, textvariable=filter_type_var, values=["All", "Income", "Expense"], state="readonly", width=10).grid(row=0, column=3, padx=3)
ttk.Label(filters, text="Category").grid(row=0, column=4, padx=3)
filter_category_combo = ttk.Combobox(filters, textvariable=filter_category_var, values=["All"], width=14)
filter_category_combo.grid(row=0, column=5, padx=3)
ttk.Label(filters, text="From").grid(row=1, column=0, padx=3)
ttk.Entry(filters, textvariable=filter_start_var, width=14).grid(row=1, column=1, padx=3)
ttk.Label(filters, text="To").grid(row=1, column=2, padx=3)
ttk.Entry(filters, textvariable=filter_end_var, width=14).grid(row=1, column=3, padx=3)
ttk.Label(filters, text="Min ₹").grid(row=1, column=4, padx=3)
ttk.Entry(filters, textvariable=min_amount_var, width=10).grid(row=1, column=5, padx=3)
ttk.Label(filters, text="Max ₹").grid(row=1, column=6, padx=3)
ttk.Entry(filters, textvariable=max_amount_var, width=10).grid(row=1, column=7, padx=3)

columns = ("ID", "Date", "Type", "Category", "Amount", "Description")
table_frame = ttk.Frame(transactions_tab)
table_frame.pack(fill="both", expand=True)
tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=12, selectmode="browse")
widths = {"ID": 55, "Date": 100, "Type": 85, "Category": 130, "Amount": 110, "Description": 350}
for col in columns:
    tree.heading(col, text=col)
    tree.column(col, width=widths[col], anchor="w" if col in ("Category", "Description", "Date", "Type") else "center")
yscroll = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
xscroll = ttk.Scrollbar(table_frame, orient="horizontal", command=tree.xview)
tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
tree.grid(row=0, column=0, sticky="nsew")
yscroll.grid(row=0, column=1, sticky="ns")
xscroll.grid(row=1, column=0, sticky="ew")
table_frame.rowconfigure(0, weight=1)
table_frame.columnconfigure(0, weight=1)

def filtered_transactions(show_errors=False):
    records = read_transactions()
    search = search_var.get().strip().lower()
    typ = filter_type_var.get()
    cat = filter_category_var.get()
    start_text, end_text = filter_start_var.get().strip(), filter_end_var.get().strip()
    min_text, max_text = min_amount_var.get().strip(), max_amount_var.get().strip()
    try:
        start = datetime.strptime(start_text, "%Y-%m-%d").date() if start_text else None
        end = datetime.strptime(end_text, "%Y-%m-%d").date() if end_text else None
        min_amt = float(min_text) if min_text else None
        max_amt = float(max_text) if max_text else None
        if min_amt is not None and min_amt < 0 or max_amt is not None and max_amt < 0:
            raise ValueError("Amount filters cannot be negative.")
        if start and end and start > end:
            raise ValueError("'From' date must not be after 'To' date.")
    except ValueError as exc:
        if show_errors:
            messagebox.showerror("Invalid filter", str(exc))
        return []
    result = []
    for r in records:
        if search and search not in " ".join(str(r[k]) for k in columns).lower():
            continue
        if typ != "All" and r["Type"] != typ:
            continue
        if cat != "All" and cat and r["Category"].lower() != cat.lower():
            continue
        try:
            d = parse_date(r["Date"])
        except ValueError:
            d = None
        if start and (d is None or d < start):
            continue
        if end and (d is None or d > end):
            continue
        if min_amt is not None and r["Amount"] < min_amt:
            continue
        if max_amt is not None and r["Amount"] > max_amt:
            continue
        result.append(r)
    return result

def display_transactions():
    for item in tree.get_children():
        tree.delete(item)
    for r in filtered_transactions():
        tree.insert("", "end", values=(r["ID"], r["Date"], r["Type"], r["Category"], money(r["Amount"]), r["Description"]))

def apply_filters():
    filtered_transactions(show_errors=True)
    display_transactions()

def clear_filters():
    for var, value in ((search_var, ""), (filter_type_var, "All"), (filter_category_var, "All"),
                       (filter_start_var, ""), (filter_end_var, ""), (min_amount_var, ""), (max_amount_var, "")):
        var.set(value)
    display_transactions()

filter_buttons = ttk.Frame(filters)
filter_buttons.grid(row=0, column=8, rowspan=2, padx=6)
ttk.Button(filter_buttons, text="Apply", command=apply_filters).pack(fill="x", pady=2)
ttk.Button(filter_buttons, text="Clear", command=clear_filters).pack(fill="x", pady=2)

def selected_record():
    selection = tree.selection()
    if not selection:
        messagebox.showwarning("Select transaction", "Select a transaction from the table first.")
        return None
    txn_id = str(tree.item(selection[0], "values")[0])
    return next((r for r in read_transactions() if str(r["ID"]) == txn_id), None)

def edit_selected():
    r = selected_record()
    if not r:
        return
    editing_id.set(str(r["ID"]))
    type_var.set(r["Type"])
    category_var.set(r["Category"])
    amount_var.set(str(r["Amount"]))
    txn_date_var.set(r["Date"])
    description_var.set(r["Description"])
    save_button.config(text="Save Changes")
    tabs.select(transactions_tab)
    category_combo.focus_set()

def delete_selected():
    r = selected_record()
    if not r:
        return
    if not messagebox.askyesno("Confirm Delete", f"Delete transaction ID {r['ID']} ({money(r['Amount'])})?"):
        return
    records = read_transactions()
    records = [x for x in records if str(x["ID"]) != str(r["ID"])]
    try:
        write_transactions(records)
        refresh_all()
        messagebox.showinfo("Deleted", "Transaction deleted.")
    except (OSError, PermissionError) as exc:
        messagebox.showerror("Delete failed", str(exc))

table_actions = ttk.Frame(transactions_tab)
table_actions.pack(fill="x", pady=6)
ttk.Button(table_actions, text="Edit Selected", command=edit_selected).pack(side="left", padx=3)
ttk.Button(table_actions, text="Delete Selected", command=delete_selected).pack(side="left", padx=3)
ttk.Button(table_actions, text="Refresh", command=lambda: refresh_all()).pack(side="left", padx=3)

# ---------------- Reports and charts ----------------
report_controls = ttk.LabelFrame(reports_tab, text="Reports", padding=8)
report_controls.pack(fill="x", pady=(0, 8))
report_period_var = tk.StringVar(value="Monthly")
report_date_var = tk.StringVar(value=date.today().isoformat())
ttk.Label(report_controls, text="Report").grid(row=0, column=0, padx=4)
ttk.Combobox(report_controls, textvariable=report_period_var,
             values=["Daily", "Weekly", "Monthly", "Yearly"], state="readonly", width=12).grid(row=0, column=1, padx=4)
ttk.Label(report_controls, text="Date in period (YYYY-MM-DD)").grid(row=0, column=2, padx=4)
ttk.Entry(report_controls, textvariable=report_date_var, width=15).grid(row=0, column=3, padx=4)
report_text = tk.Text(reports_tab, height=10, wrap="word")
report_text.pack(fill="x", pady=5)
chart_frame = ttk.Frame(reports_tab)
chart_frame.pack(fill="both", expand=True)

def report_range():
    d = datetime.strptime(report_date_var.get().strip(), "%Y-%m-%d").date()
    period = report_period_var.get()
    if period == "Daily":
        return d, d, d.strftime("%Y-%m-%d")
    if period == "Weekly":
        start = d - timedelta(days=d.weekday())
        return start, start + timedelta(days=6), f"Week {start.isoformat()} to {(start + timedelta(days=6)).isoformat()}"
    if period == "Monthly":
        start = date(d.year, d.month, 1)
        if d.month == 12:
            next_month = date(d.year + 1, 1, 1)
        else:
            next_month = date(d.year, d.month + 1, 1)
        return start, next_month - timedelta(days=1), d.strftime("%B %Y")
    return date(d.year, 1, 1), date(d.year, 12, 31), str(d.year)

def generate_report():
    try:
        start, end, label = report_range()
    except ValueError:
        messagebox.showerror("Invalid date", "Enter the report date as YYYY-MM-DD.")
        return
    records = period_records(read_transactions(), start, end)
    inc, exp, balance = totals(records)
    by_cat = defaultdict(float)
    for r in records:
        if r["Type"] == "Expense":
            by_cat[r["Category"]] += r["Amount"]
    lines = [f"{report_period_var.get()} Report: {label}", f"Period: {start} to {end}",
             f"Total income: {money(inc)}", f"Total expenses: {money(exp)}",
             f"Net savings/balance: {money(balance)}", "", "Expense by category:"]
    lines += [f"  • {cat}: {money(amount)}" for cat, amount in sorted(by_cat.items(), key=lambda x: x[1], reverse=True)]
    if not by_cat:
        lines.append("  No expenses recorded in this period.")
    report_text.delete("1.0", tk.END)
    report_text.insert("1.0", "\n".join(lines))

def clear_chart():
    for child in chart_frame.winfo_children():
        child.destroy()

def show_category_chart():
    records = [r for r in read_transactions() if r["Type"] == "Expense"]
    by_cat = defaultdict(float)
    for r in records:
        by_cat[r["Category"]] += r["Amount"]
    if not by_cat:
        messagebox.showinfo("No data", "Add expenses before creating a chart.")
        return
    clear_chart()
    fig, ax = plt.subplots(figsize=(7, 4.2))
    ax.pie(list(by_cat.values()), labels=list(by_cat.keys()), autopct="%1.1f%%", startangle=90)
    ax.set_title("Expenses by Category (All Records)")
    fig.tight_layout()
    canvas = FigureCanvasTkAgg(fig, master=chart_frame)
    canvas.draw()
    canvas.get_tk_widget().pack(fill="both", expand=True)
    chart_frame._figure = fig

def show_monthly_comparison():
    records = read_transactions()
    month_totals = defaultdict(lambda: {"Income": 0.0, "Expense": 0.0})
    for r in records:
        try:
            key = parse_date(r["Date"]).strftime("%Y-%m")
        except ValueError:
            continue
        month_totals[key][r["Type"]] += r["Amount"]
    if not month_totals:
        messagebox.showinfo("No data", "No transactions available for comparison.")
        return
    clear_chart()
    months = sorted(month_totals)[-12:]
    fig, ax = plt.subplots(figsize=(8, 4.2))
    x = list(range(len(months)))
    width = 0.38
    ax.bar([i - width/2 for i in x], [month_totals[m]["Income"] for m in months], width, label="Income")
    ax.bar([i + width/2 for i in x], [month_totals[m]["Expense"] for m in months], width, label="Expenses")
    ax.set_xticks(x)
    ax.set_xticklabels(months, rotation=45, ha="right")
    ax.set_ylabel("Amount (₹)")
    ax.set_title("Monthly Income vs Expenses")
    ax.legend()
    fig.tight_layout()
    canvas = FigureCanvasTkAgg(fig, master=chart_frame)
    canvas.draw()
    canvas.get_tk_widget().pack(fill="both", expand=True)
    chart_frame._figure = fig

report_buttons = ttk.Frame(reports_tab)
report_buttons.pack(fill="x", pady=5)
ttk.Button(report_buttons, text="Generate Report", command=generate_report).pack(side="left", padx=3)
ttk.Button(report_buttons, text="Category Pie Chart", command=show_category_chart).pack(side="left", padx=3)
ttk.Button(report_buttons, text="Monthly Comparison", command=show_monthly_comparison).pack(side="left", padx=3)

# ---------------- Budget management ----------------
budget_form = ttk.LabelFrame(budgets_tab, text="Set Monthly Category Budget", padding=8)
budget_form.pack(fill="x", pady=(0, 8))
budget_category_var = tk.StringVar(value="Food")
budget_amount_var = tk.StringVar()
ttk.Label(budget_form, text="Category").grid(row=0, column=0, padx=4, pady=4)
budget_category_combo = ttk.Combobox(budget_form, textvariable=budget_category_var, values=DEFAULT_CATEGORIES, width=20)
budget_category_combo.grid(row=0, column=1, padx=4, pady=4)
ttk.Label(budget_form, text="Budget amount (₹)").grid(row=0, column=2, padx=4, pady=4)
ttk.Entry(budget_form, textvariable=budget_amount_var, width=15).grid(row=0, column=3, padx=4, pady=4)
budget_tree = ttk.Treeview(budgets_tab, columns=("Category", "Budget", "Spent", "Remaining", "Status"), show="headings", height=12)
for c, w in (("Category", 180), ("Budget", 140), ("Spent", 140), ("Remaining", 140), ("Status", 180)):
    budget_tree.heading(c, text=c)
    budget_tree.column(c, width=w, anchor="center")
budget_tree.pack(fill="both", expand=True)

def refresh_budgets():
    for item in budget_tree.get_children():
        budget_tree.delete(item)
    budgets = load_json(BUDGET_FILE, {})
    start, end = current_month_bounds()
    records = period_records(read_transactions(), start, end)
    spent = defaultdict(float)
    for r in records:
        if r["Type"] == "Expense":
            spent[r["Category"]] += r["Amount"]
    for cat, limit in sorted(budgets.items(), key=lambda x: x[0].lower()):
        limit = float(limit)
        amount = spent[cat]
        remaining = limit - amount
        status = "OVER BUDGET" if amount > limit else ("Limit reached" if amount >= limit else "Within budget")
        budget_tree.insert("", "end", values=(cat, money(limit), money(amount), money(remaining), status))

def save_budget():
    cat = budget_category_var.get().strip()
    try:
        amount = float(budget_amount_var.get().strip())
        if not cat or amount <= 0 or amount == float("inf"):
            raise ValueError
    except ValueError:
        messagebox.showerror("Invalid budget", "Enter a category and a budget amount greater than zero.")
        return
    budgets = load_json(BUDGET_FILE, {})
    budgets[cat] = round(amount, 2)
    save_json(BUDGET_FILE, budgets)
    budget_amount_var.set("")
    refresh_budgets()
    check_budget_alerts()
    messagebox.showinfo("Budget saved", f"Monthly budget for {cat} saved.")

def delete_budget():
    selection = budget_tree.selection()
    if not selection:
        messagebox.showwarning("Select budget", "Select a budget row first.")
        return
    cat = budget_tree.item(selection[0], "values")[0]
    budgets = load_json(BUDGET_FILE, {})
    budgets.pop(cat, None)
    save_json(BUDGET_FILE, budgets)
    refresh_budgets()

ttk.Button(budget_form, text="Save Budget", command=save_budget).grid(row=0, column=4, padx=5)
ttk.Button(budget_form, text="Delete Selected Budget", command=delete_budget).grid(row=0, column=5, padx=5)

def check_budget_alerts():
    budgets = load_json(BUDGET_FILE, {})
    start, end = current_month_bounds()
    spent = defaultdict(float)
    for r in period_records(read_transactions(), start, end):
        if r["Type"] == "Expense":
            spent[r["Category"]] += r["Amount"]
    over = [f"{cat}: {money(spent[cat])} spent of {money(float(limit))}"
            for cat, limit in budgets.items() if spent[cat] >= float(limit)]
    if over:
        insights_var.set(insights_var.get() + "  |  Budget alert: " + "; ".join(over))

# ---------------- Recurring transactions ----------------
rec_form = ttk.LabelFrame(recurring_tab, text="Set up a recurring transaction", padding=8)
rec_form.pack(fill="x", pady=(0, 8))
rec_type_var = tk.StringVar(value="Expense")
rec_cat_var = tk.StringVar(value="Rent")
rec_amt_var = tk.StringVar()
rec_desc_var = tk.StringVar()
rec_day_var = tk.StringVar(value=str(min(date.today().day, 28)))
ttk.Label(rec_form, text="Type").grid(row=0, column=0, padx=3, pady=4)
ttk.Combobox(rec_form, textvariable=rec_type_var, values=["Income", "Expense"], state="readonly", width=10).grid(row=0, column=1, padx=3)
ttk.Label(rec_form, text="Category").grid(row=0, column=2, padx=3)
ttk.Entry(rec_form, textvariable=rec_cat_var, width=15).grid(row=0, column=3, padx=3)
ttk.Label(rec_form, text="Amount ₹").grid(row=0, column=4, padx=3)
ttk.Entry(rec_form, textvariable=rec_amt_var, width=12).grid(row=0, column=5, padx=3)
ttk.Label(rec_form, text="Day (1-28)").grid(row=1, column=0, padx=3, pady=4)
ttk.Entry(rec_form, textvariable=rec_day_var, width=8).grid(row=1, column=1, padx=3)
ttk.Label(rec_form, text="Description").grid(row=1, column=2, padx=3)
ttk.Entry(rec_form, textvariable=rec_desc_var, width=32).grid(row=1, column=3, columnspan=2, padx=3)
rec_tree = ttk.Treeview(recurring_tab, columns=("ID", "Type", "Category", "Amount", "Day", "Description"), show="headings", height=12)
for c, w in (("ID", 55), ("Type", 90), ("Category", 140), ("Amount", 120), ("Day", 80), ("Description", 300)):
    rec_tree.heading(c, text=c)
    rec_tree.column(c, width=w, anchor="center")
rec_tree.pack(fill="both", expand=True)

def refresh_recurring():
    for item in rec_tree.get_children():
        rec_tree.delete(item)
    for i, r in enumerate(load_json(RECURRING_FILE, []), start=1):
        rec_tree.insert("", "end", iid=str(r.get("id", i)), values=(r.get("id", i), r["Type"], r["Category"], money(float(r["Amount"])), r["Day"], r.get("Description", "")))

def add_recurring():
    cat = rec_cat_var.get().strip()
    try:
        amount = float(rec_amt_var.get().strip())
        day = int(rec_day_var.get().strip())
        if not cat or amount <= 0 or amount == float("inf") or not 1 <= day <= 28:
            raise ValueError
        if rec_type_var.get() not in ("Income", "Expense"):
            raise ValueError
    except ValueError:
        messagebox.showerror("Invalid recurring transaction", "Enter category, valid amount, and day from 1 to 28.")
        return
    items = load_json(RECURRING_FILE, [])
    new_id = max([int(x.get("id", 0)) for x in items] or [0]) + 1
    items.append({"id": new_id, "Type": rec_type_var.get(), "Category": cat, "Amount": round(amount, 2),
                  "Day": day, "Description": rec_desc_var.get().strip(), "last_created": ""})
    save_json(RECURRING_FILE, items)
    rec_amt_var.set("")
    refresh_recurring()
    messagebox.showinfo("Recurring saved", "Recurring rule saved. It will not add duplicate transactions automatically.")

def create_due_recurring():
    """Manually create the recurring items for the current month, at most once per rule/month."""
    today = date.today()
    items = load_json(RECURRING_FILE, [])
    records = read_transactions()
    added = 0
    for item in items:
        month_key = today.strftime("%Y-%m")
        if item.get("last_created") == month_key or today.day < int(item["Day"]):
            continue
        txn_date = date(today.year, today.month, int(item["Day"])).isoformat()
        records.append({"ID": next_id(records), "Date": txn_date, "Type": item["Type"],
                        "Category": item["Category"], "Amount": float(item["Amount"]),
                        "Description": item.get("Description", "Recurring transaction")})
        item["last_created"] = month_key
        added += 1
    if added:
        write_transactions(records)
        save_json(RECURRING_FILE, items)
        refresh_all()
    messagebox.showinfo("Recurring transactions", f"Created {added} transaction(s) for this month.")

def delete_recurring():
    selection = rec_tree.selection()
    if not selection:
        messagebox.showwarning("Select recurring rule", "Select a recurring rule first.")
        return
    rid = str(rec_tree.item(selection[0], "values")[0])
    items = [x for x in load_json(RECURRING_FILE, []) if str(x.get("id")) != rid]
    save_json(RECURRING_FILE, items)
    refresh_recurring()

rec_buttons = ttk.Frame(recurring_tab)
rec_buttons.pack(fill="x", pady=6)
ttk.Button(rec_form, text="Save Recurring Rule", command=add_recurring).grid(row=1, column=5, padx=4)
ttk.Button(rec_buttons, text="Create Due Items This Month", command=create_due_recurring).pack(side="left", padx=3)
ttk.Button(rec_buttons, text="Delete Selected Rule", command=delete_recurring).pack(side="left", padx=3)
ttk.Label(recurring_tab, text="Recurring entries are created only when you press the button; this avoids silent/duplicate entries.", wraplength=900).pack(anchor="w", pady=4)

# ---------------- Export, backup, restore ----------------
ttk.Label(tools_tab, text="Export and data safety", font=("Arial", 13, "bold")).pack(anchor="w", pady=5)
ttk.Label(tools_tab, text="Excel export saves a separate copy. Backup copies your current workbook. Restore replaces the current workbook only after confirmation.", wraplength=900).pack(anchor="w", pady=5)

def export_csv():
    path = filedialog.asksaveasfilename(title="Export transactions to CSV", defaultextension=".csv",
                                        filetypes=[("CSV files", "*.csv")], initialfile="expenses_export.csv")
    if not path:
        return
    try:
        records = read_transactions()
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=HEADERS)
            writer.writeheader()
            writer.writerows(records)
        messagebox.showinfo("Export complete", f"CSV exported to:\n{path}")
    except (OSError, PermissionError) as exc:
        messagebox.showerror("Export failed", str(exc))

def export_excel():
    path = filedialog.asksaveasfilename(title="Export transactions to Excel", defaultextension=".xlsx",
                                        filetypes=[("Excel files", "*.xlsx")], initialfile="expenses_export.xlsx")
    if not path:
        return
    try:
        records = read_transactions()
        wb = Workbook()
        ws = wb.active
        ws.title = "Transactions"
        ws.append(HEADERS)
        for r in records:
            ws.append([r["ID"], r["Date"], r["Type"], r["Category"], r["Amount"], r["Description"]])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            letter = column[0].column_letter
            ws.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in column) + 2, 12), 36)
        wb.save(path)
        wb.close()
        messagebox.showinfo("Export complete", f"Excel exported to:\n{path}")
    except (OSError, PermissionError) as exc:
        messagebox.showerror("Export failed", str(exc))

def backup_database():
    if not EXCEL_FILE.exists():
        messagebox.showerror("No database", "Excel database file was not found.")
        return
    backup_dir = DATA_DIR / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    destination = backup_dir / f"expenses_backup_{stamp}.xlsx"
    try:
        shutil.copy2(EXCEL_FILE, destination)
        messagebox.showinfo("Backup complete", f"Backup saved:\n{destination}")
    except OSError as exc:
        messagebox.showerror("Backup failed", str(exc))

def restore_database():
    source = filedialog.askopenfilename(title="Select Excel backup to restore",
                                        filetypes=[("Excel files", "*.xlsx")])
    if not source:
        return
    source_path = Path(source).resolve()
    if source_path == EXCEL_FILE.resolve():
        messagebox.showwarning("Same file", "Choose a backup file, not the current database.")
        return
    try:
        check = load_workbook(source_path, read_only=True, data_only=True)
        valid = "Transactions" in check.sheetnames and [check["Transactions"].cell(1, c).value for c in range(1, 7)] == HEADERS
        check.close()
        if not valid:
            messagebox.showerror("Invalid backup", "Selected file does not have the expected Transactions sheet and headers.")
            return
    except Exception as exc:
        messagebox.showerror("Invalid backup", f"Could not read the selected workbook:\n{exc}")
        return
    if not messagebox.askyesno("Confirm restore", "Restore will replace the current Excel workbook. A safety backup will be made first. Continue?"):
        return
    try:
        backup_database_silent()
        shutil.copy2(source_path, EXCEL_FILE)
        refresh_all()
        messagebox.showinfo("Restore complete", "Database restored successfully.")
    except OSError as exc:
        messagebox.showerror("Restore failed", str(exc))

def backup_database_silent():
    backup_dir = DATA_DIR / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(EXCEL_FILE, backup_dir / f"before_restore_{stamp}.xlsx")

def open_data_folder():
    import os
    try:
        os.startfile(str(DATA_DIR))
    except OSError as exc:
        messagebox.showerror("Cannot open folder", str(exc))

tool_buttons = ttk.Frame(tools_tab)
tool_buttons.pack(anchor="w", pady=8)
ttk.Button(tool_buttons, text="Export CSV", command=export_csv).grid(row=0, column=0, padx=4, pady=4)
ttk.Button(tool_buttons, text="Export Excel", command=export_excel).grid(row=0, column=1, padx=4, pady=4)
ttk.Button(tool_buttons, text="Backup Database", command=backup_database).grid(row=0, column=2, padx=4, pady=4)
ttk.Button(tool_buttons, text="Restore Backup", command=restore_database).grid(row=0, column=3, padx=4, pady=4)
ttk.Button(tool_buttons, text="Open Data Folder", command=open_data_folder).grid(row=0, column=4, padx=4, pady=4)

# ---------------- Dashboard calculations and refresh ----------------
def refresh_dashboard():
    records = read_transactions()
    income, expenses, balance = totals(records)
    income_label.config(text=f"Total Income: {money(income)}")
    expense_label.config(text=f"Total Expenses: {money(expenses)}")
    balance_label.config(text=f"Current Balance: {money(balance)}")
    savings_rate = ((balance / income) * 100) if income else None
    savings_label.config(text=f"Savings Rate: {savings_rate:.1f}%" if savings_rate is not None else "Savings Rate: N/A")
    expense_records = [r for r in records if r["Type"] == "Expense"]
    by_cat = defaultdict(float)
    for r in expense_records:
        by_cat[r["Category"]] += r["Amount"]
    highest = max(by_cat, key=by_cat.get) if by_cat else None
    days = set()
    for r in expense_records:
        try:
            days.add(parse_date(r["Date"]))
        except ValueError:
            pass
    average = sum(r["Amount"] for r in expense_records) / len(days) if days else 0
    alert = "Expenses exceed income." if expenses > income else "Expenses are within recorded income."
    parts = [
        f"Highest spending category: {highest} ({money(by_cat[highest])})" if highest else "Highest spending category: no expenses yet",
        f"Average expense per day with expenses: {money(average)}",
        alert
    ]
    insights_var.set("  |  ".join(parts))
    refresh_category_choices()
    filter_category_combo["values"] = ["All"] + get_categories()
    budget_category_combo["values"] = get_categories()
    display_transactions()

def refresh_all():
    refresh_dashboard()
    refresh_budgets()
    refresh_recurring()

def on_close():
    for child in chart_frame.winfo_children():
        child.destroy()
    plt.close("all")
    root.destroy()

root.protocol("WM_DELETE_WINDOW", on_close)

try:
    refresh_all()
except Exception as exc:
    messagebox.showerror("Startup error", f"Could not load the expense database:\n{exc}")

root.mainloop()
